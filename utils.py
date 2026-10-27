"""Helpers: login guard, receipt/report export (PDF + Excel), QR image finder."""
import io
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

BASE = Path(__file__).parent
SHOP = "Makara Café"


# ---------------- login / role ----------------
def require_login(roles=None):
    """Call at the top of every view. Stops the page if not logged in or role not allowed."""
    if "user" not in st.session_state:
        st.warning("Please log in first.")
        st.stop()
    user = st.session_state["user"]
    if roles and user["role"] not in roles:
        st.error("⛔ You don't have permission to open this page.")
        st.stop()
    return user


def qr_png(text):
    """Turn a text payload (e.g. KHQR string) into PNG bytes for st.image."""
    import qrcode
    img = qrcode.make(text, box_size=8, border=2)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


FA_CSS = "<style>@import url('https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.5.2/css/all.min.css');</style>"


def html(markup):
    """Show raw HTML/CSS. Uses st.markdown(unsafe_allow_html) because it works in every Streamlit version."""
    one_line = " ".join(line.strip() for line in markup.strip().splitlines())  # no blank lines / indents
    st.markdown(one_line, unsafe_allow_html=True)


def load_icons():
    """Load Font Awesome (needs internet). Called once in app.py."""
    html(FA_CSS)


def load_css():
    """Inject style.css (edit that file to change colours, cards, spacing...)."""
    css = (BASE / "style.css").read_text(encoding="utf-8")
    html(f"<style>{css}</style>")


def fa(name, size=24, color="#6F4E37"):
    """HTML for a Font Awesome solid icon, e.g. fa('mug-hot')."""
    return f'<i class="fa-solid fa-{name}" style="font-size:{size}px;color:{color}"></i>'


def title(icon, text, color="#6F4E37"):
    """Page title with a Font Awesome icon."""
    from html import escape
    html(f'<h1 style="font-size:2.25rem;font-weight:700;margin:0 0 .75rem 0;padding:0">'
            f'{fa(icon, 32, color)}&nbsp; {escape(text)}</h1>')


def flash(msg):
    st.session_state["_flash"] = msg


def show_flash():
    msg = st.session_state.pop("_flash", None)
    if msg:
        st.success(msg)


IMAGES = BASE / "static" / "products"   # served by Streamlit at /app/static/products/<file>
OLD_IMAGES = BASE / "images"


def save_product_image(file):
    """Resize + crop an uploaded photo to 4:3 so every menu card looks the same. Returns the file name."""
    import uuid
    from PIL import Image, ImageOps
    IMAGES.mkdir(parents=True, exist_ok=True)
    img = ImageOps.exif_transpose(Image.open(file)).convert("RGB")
    img = ImageOps.fit(img, (480, 360), Image.LANCZOS)
    name = f"{uuid.uuid4().hex[:12]}.jpg"
    img.save(IMAGES / name, "JPEG", quality=85)
    return name


def save_image_from_url(url, allow_private=False):
    """Download a picture from a direct image link (right-click the image -> Copy image address),
    resize it like an upload and save it. Returns the file name."""
    import ipaddress, socket
    from urllib.parse import urlparse
    from urllib.request import Request, urlopen
    u = urlparse((url or "").strip())
    if u.scheme not in ("http", "https") or not u.hostname:
        raise ValueError("Please paste a full link starting with http:// or https://")
    if not allow_private:  # do not let the app fetch things from the local network
        for info in socket.getaddrinfo(u.hostname, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local:
                raise ValueError("That address is not allowed")
    req = Request(url.strip(), headers={"User-Agent": "Mozilla/5.0 (Makara Cafe POS)"})
    try:
        with urlopen(req, timeout=12) as r:
            data = r.read(8_000_001)
    except Exception as e:
        raise ValueError(f"Could not download the picture ({e.__class__.__name__}). Use a direct image link.")
    if len(data) > 8_000_000:
        raise ValueError("Picture is too big (max 8 MB)")
    try:
        return save_product_image(io.BytesIO(data))
    except Exception:
        raise ValueError("That link is not a picture. Use a direct image link (right-click the image -> Copy image address).")


def product_image(name):
    """Path of a product photo, or None if there is none."""
    if not isinstance(name, str) or not name:
        return None
    p = IMAGES / name
    old = OLD_IMAGES / name
    if not p.exists() and old.exists():  # photos saved by the previous version -> move to static/
        IMAGES.mkdir(parents=True, exist_ok=True)
        p.write_bytes(old.read_bytes())
    return p if p.exists() else None


def slug(text):
    import re
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")


def photo_for(product):
    """Photo file name for a product: the uploaded one, or a file in static/products named after the product
    (e.g. 'iced-coffee.jpg' for 'Iced Coffee'). Returns None if there is no real photo."""
    if product_image(product.get("image")):
        return product["image"]
    if IMAGES.exists():
        want = slug(product["name"])
        for f in sorted(IMAGES.iterdir()):
            if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp") and slug(f.stem) == want:
                return f.name
    return None


def assign_photos(files):
    """Bulk upload: match each file to a product by file name ('iced coffee.jpg' -> Iced Coffee).
    Returns (matched, unmatched) lists of file names."""
    import db
    prods = {slug(r["name"]): r for r in db.list_products().to_dict("records")}
    matched, unmatched = [], []
    for f in files:
        key = slug(f.name.rsplit(".", 1)[0])
        p = prods.get(key) or next((r for k, r in prods.items() if key and (key in k or k in key)), None)
        if p is None:
            unmatched.append(f.name)
            continue
        db.update_product(int(p["id"]), p["name"], p["price"], p["category"], save_product_image(f))
        matched.append(f"{f.name} → {p['name']}")
    return matched, unmatched


def find_qr_image():
    """First image in Photo/ (prefers a file with 'qr' in its name)."""
    files = [f for f in sorted((BASE / "Photo").glob("*")) if f.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp")]
    for f in files:
        if "qr" in f.name.lower():
            return f
    return files[0] if files else None


# ---------------- PDF helpers ----------------
def _s(x):
    """fpdf core fonts are Latin-1 only, so replace unsupported characters."""
    if isinstance(x, float):
        x = f"{x:,.2f}"
    return str(x).encode("latin-1", "replace").decode("latin-1")


def receipt_pdf(order, items):
    h = 70 + 7 * len(items)
    pdf = FPDF(format=(80, h))
    pdf.set_margins(5, 5, 5)
    pdf.set_auto_page_break(False)
    pdf.add_page()
    nl = dict(new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(70, 8, SHOP, align="C", **nl)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(70, 5, "Phnom Penh, Cambodia", align="C", **nl)
    pdf.cell(70, 5, f"Receipt #{order['id']}", align="C", **nl)
    pdf.cell(70, 5, _s(order["date"]), align="C", **nl)
    pdf.cell(70, 5, f"Cashier: {_s(order['cashier'])}", align="C", **nl)
    pdf.cell(70, 4, "-" * 52, align="C", **nl)
    pdf.set_font("Helvetica", "B", 8)
    pdf.cell(35, 5, "Item"); pdf.cell(10, 5, "Qty", align="R"); pdf.cell(25, 5, "Amount", align="R", **nl)
    pdf.set_font("Helvetica", "", 8)
    for _, r in items.iterrows():
        pdf.cell(35, 5, _s(r["Product"])[:22]); pdf.cell(10, 5, str(int(r["Qty"])), align="R")
        pdf.cell(25, 5, f"${r['Amount']:.2f}", align="R", **nl)
    pdf.cell(70, 4, "-" * 52, align="C", **nl)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(45, 7, "TOTAL"); pdf.cell(25, 7, f"${order['total']:.2f}", align="R", **nl)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(70, 5, f"Payment: {_s(order['payment_method'])}", **nl)
    pdf.ln(2)
    pdf.cell(70, 5, "Thank you! See you again.", align="C", **nl)
    return bytes(pdf.output())


def receipt_excel(order, items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Receipt"
    ws.append([SHOP]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Receipt #{order['id']}"])
    ws.append(["Date", order["date"]])
    ws.append(["Cashier", order["cashier"]])
    ws.append(["Payment", order["payment_method"]])
    ws.append([])
    ws.append(["Product", "Qty", "Price", "Amount"])
    for c in ws[7]:
        c.font = Font(bold=True)
    for _, r in items.iterrows():
        ws.append([r["Product"], int(r["Qty"]), float(r["Price"]), float(r["Amount"])])
    ws.append(["TOTAL", "", "", float(order["total"])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, w in zip("ABCD", (28, 8, 10, 12)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_pdf(title, subtitle, sections):
    """sections = [(heading, DataFrame), ...]"""
    pdf = FPDF()
    pdf.add_page()
    nl = dict(new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, _s(f"{SHOP} - {title}"), **nl)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _s(subtitle), **nl)
    pdf.cell(0, 6, f"Generated: {datetime.now():%Y-%m-%d %H:%M}", **nl)
    pdf.ln(4)
    for heading, df in sections:
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, _s(heading), **nl)
        if df.empty:
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 6, "No data", **nl)
            continue
        w = 190 / len(df.columns)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(230, 230, 230)
        for c in df.columns:
            pdf.cell(w, 7, _s(c), border=1, fill=True)
        pdf.ln()
        pdf.set_font("Helvetica", "", 9)
        for _, row in df.iterrows():
            if pdf.get_y() > 275:
                pdf.add_page()
            for v in row:
                pdf.cell(w, 6, _s(v)[:28], border=1)
            pdf.ln()
        pdf.ln(4)
    return bytes(pdf.output())


def report_excel(sheets):
    """sheets = {sheet_name: DataFrame}"""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for i, col in enumerate(df.columns, start=1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(str(col)) + 4)
    return buf.getvalue()
